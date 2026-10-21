"""Extract auditable, aggregated reporting data without modifying source workbooks.

Run with the bundled Python runtime. Requires openpyxl for read-only extraction.
No employee names, merchants or transaction-level personal data are exported.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "LATAM (Jan-Sep).xlsx"
OUTPUT = HERE / "financial-data.json"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September"]
COUNTRIES = [("AR", "Argentina", 5), ("BR", "Brazil", 18), ("CL", "Chile", 31),
             ("CO", "Colombia", 44), ("MX", "Mexico", 57), ("PE", "Peru", 70)]
FIELDS = ["income", "billables", "managementFees", "revenue", "expenses", "profit"]
COSTS = ["card", "payroll", "pmSalary", "seatCharges", "procurement", "cashTransfers", "otherIndirect"]


def num(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def tidy(value):
    if isinstance(value, float):
        return round(value, 8)
    if isinstance(value, dict):
        return {key: tidy(item) for key, item in value.items()}
    if isinstance(value, list):
        return [tidy(item) for item in value]
    return value


def labels(pt, es, en):
    return {"pt": pt, "es": es, "en": en}


def main():
    values = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
    pl = values["P&L"]
    expenses = values["Expenses"]
    raw = list(values["Card Raw"].iter_rows(values_only=True))
    assert raw[0] == ("Transaction Date", "User", "Merchant Name", "Ramp Category", "Ramp Department", "Statement Amount", "Month")
    dated = []
    for row in raw[1:]:
        dt = row[0]
        dated.append(dt if isinstance(dt, datetime) else datetime.strptime(dt, "%m/%d/%Y"))
    years = sorted({dt.year for dt in dated})
    assert len(years) == 1
    year = years[0]
    assert all(MONTHS[dt.month - 1] == row[6] for dt, row in zip(dated, raw[1:]))

    records = []
    component_checks = []
    monthly_checks = []
    country_checks = []
    missing_components = []
    for country_index, (code, country, start_row) in enumerate(COUNTRIES):
        assert pl.cell(start_row - 1, 1).value == country
        for month_index, month_name in enumerate(MONTHS):
            month = month_index + 1
            row = start_row + month_index
            assert pl.cell(row, 2).value == month_name
            record = {"country": country, "countryCode": code, "month": month,
                      "period": f"{year}-{month:02d}"}
            for col, key in enumerate(FIELDS, start=3):
                value = pl.cell(row, col).value
                if value is None and key == "billables":
                    # The workbook explicitly includes this blank in C+D+E and caches F=C.
                    assert pl.cell(row, 6).value == pl.cell(row, 3).value
                    value = 0
                    missing_components.append({"countryCode": code, "month": month, "field": key,
                                               "cell": f"P&L!{get_column_letter(col)}{row}",
                                               "treatment": "zero_as_used_by_source_formula"})
                assert num(value), (country, month, key, value)
                record[key] = value
            record["margin"] = record["profit"] / record["revenue"] if record["revenue"] else None
            record["returnOnCost"] = record["profit"] / record["expenses"] if record["expenses"] else None
            card_ref = f"{get_column_letter(month+1)}{5+country_index}"
            payroll_ref = f"{get_column_letter(month+1)}{19+country_index}"
            pm_ref = f"{get_column_letter(month+1)}29"
            seat_ref = f"{get_column_letter(month+1)}30"
            procurement_ref = f"AG{month+2}"
            cash_ref = f"{get_column_letter(month+18)}{4+country_index}"
            indirect_ref = f"AF{month+2}"
            cost_refs = dict(zip(COSTS, [card_ref, payroll_ref, pm_ref, seat_ref, procurement_ref, cash_ref, indirect_ref]))
            record["costs"] = {key: expenses[ref].value if expenses[ref].value is not None else 0
                               for key, ref in cost_refs.items()}
            record["allocatedOverhead"] = sum(record["costs"][key] for key in ["pmSalary", "seatCharges", "procurement", "otherIndirect"])
            record["localExpenses"] = record["expenses"] - record["allocatedOverhead"]
            record["contributionProfit"] = record["revenue"] - record["localExpenses"]
            record["contributionMargin"] = record["contributionProfit"] / record["revenue"] if record["revenue"] else None
            component_checks.append(record["expenses"] - sum(record["costs"].values()))
            assert abs(component_checks[-1]) < .00001, (country, month, component_checks[-1])
            assert abs(record["income"] + record["billables"] + record["managementFees"] - record["revenue"]) < .00001
            assert abs(record["revenue"] - record["expenses"] - record["profit"]) < .00001
            assert abs(record["billables"] * .078 - record["managementFees"]) < .00001
            record["source"] = {"sheet": "P&L", "range": f"B{row}:H{row}",
                                "costSheet": "Expenses", "costCells": cost_refs}
            records.append(record)
        for col, key in enumerate(FIELDS, start=3):
            original = pl.cell(start_row+9, col).value
            derived = sum(record[key] for record in records if record["countryCode"] == code)
            country_checks.append(derived-original)
            assert abs(derived-original) < .00001

    def aggregate(subset):
        result = {key: sum(row[key] for row in subset) for key in FIELDS}
        result["margin"] = result["profit"] / result["revenue"] if result["revenue"] else None
        result["returnOnCost"] = result["profit"] / result["expenses"] if result["expenses"] else None
        result["costs"] = {key: sum(row["costs"][key] for row in subset) for key in COSTS}
        for key in ["allocatedOverhead", "localExpenses", "contributionProfit"]:
            result[key] = sum(row[key] for row in subset)
        result["contributionMargin"] = result["contributionProfit"] / result["revenue"] if result["revenue"] else None
        return result

    totals = aggregate(records)
    monthly = [{"month": month, "period": f"{year}-{month:02d}",
                **aggregate([record for record in records if record["month"] == month])}
               for month in range(1, 10)]
    country_totals = [{"countryCode": code, "country": country,
                       **aggregate([record for record in records if record["countryCode"] == code])}
                      for code, country, _ in COUNTRIES]

    # Merge only spelling/wording variants. The source category remains traceable.
    aliases = {
        "Fuel & Gas": "Fuel and Gas", "Restaurant": "Restaurants",
        "Hotel & Lodging": "Lodging", "Car services": "Car Services", "Car Service": "Car Services",
        "Taxi & Rideshare": "Taxi and Rideshare", "Internet & Phone": "Internet and Phone",
        "Phone and Internet": "Internet and Phone", "Office and Supplies": "Office Supplies",
    }
    category_labels = {
        "Advertising": labels("Publicidade", "Publicidad", "Advertising"),
        "Airlines": labels("Passagens aéreas", "Pasajes aéreos", "Airlines"),
        "Alcohol and Bars": labels("Bares e bebidas", "Bares y bebidas", "Alcohol and bars"),
        "Baggage": labels("Bagagem", "Equipaje", "Baggage"),
        "Books and Newspapers": labels("Livros e jornais", "Libros y periódicos", "Books and newspapers"),
        "Car Rental": labels("Aluguel de veículos", "Alquiler de vehículos", "Car rental"),
        "Car Services": labels("Serviços de veículos", "Servicios de vehículos", "Vehicle services"),
        "Charitable donations": labels("Doações", "Donaciones", "Donations"),
        "Clothing": labels("Vestuário", "Ropa", "Clothing"),
        "Clubs and memberships": labels("Clubes e associações", "Clubes y membresías", "Clubs and memberships"),
        "Education": labels("Educação", "Educación", "Education"),
        "Electronics": labels("Eletrônicos", "Electrónica", "Electronics"),
        "Entertainment": labels("Entretenimento", "Entretenimiento", "Entertainment"),
        "Fees and Financial institutions": labels("Tarifas financeiras", "Comisiones financieras", "Financial fees"),
        "Freight, Moving and Delivery Services": labels("Fretes e entregas", "Fletes y entregas", "Freight and delivery"),
        "Fuel and Gas": labels("Combustível", "Combustible", "Fuel and gas"),
        "General Merchandise": labels("Mercadorias em geral", "Mercancías generales", "General merchandise"),
        "Government Services": labels("Serviços públicos", "Servicios gubernamentales", "Government services"),
        "Insurance": labels("Seguros", "Seguros", "Insurance"),
        "Internet and Phone": labels("Internet e telefone", "Internet y teléfono", "Internet and phone"),
        "Laundry": labels("Lavanderia", "Lavandería", "Laundry"),
        "Lodging": labels("Hospedagem", "Alojamiento", "Lodging"),
        "Medical": labels("Saúde", "Salud", "Medical"),
        "Office": labels("Escritório", "Oficina", "Office"),
        "Office Supplies": labels("Materiais de escritório", "Material de oficina", "Office supplies"),
        "Office supplies and cleaning": labels("Escritório e limpeza", "Oficina y limpieza", "Office supplies and cleaning"),
        "Other": labels("Outros", "Otros", "Other"),
        "Parking": labels("Estacionamento", "Estacionamiento", "Parking"),
        "Pet": labels("Animais", "Animales", "Pet"),
        "Professional Services": labels("Serviços profissionais", "Servicios profesionales", "Professional services"),
        "Restaurants": labels("Restaurantes", "Restaurantes", "Restaurants"),
        "SaaS / Software": labels("Software e SaaS", "Software y SaaS", "Software and SaaS"),
        "Shipping": labels("Envios", "Envíos", "Shipping"),
        "Streaming Services": labels("Streaming", "Streaming", "Streaming services"),
        "Supermarkets and Grocery Stores": labels("Supermercados", "Supermercados", "Supermarkets and groceries"),
        "Tag": labels("Tag", "Tag", "Tag"),
        "Taxes and tax preparation": labels("Impostos", "Impuestos", "Taxes and tax preparation"),
        "Taxi and Rideshare": labels("Táxi e aplicativos", "Taxi y aplicaciones", "Taxi and rideshare"),
        "Toll": labels("Pedágios", "Peajes", "Tolls"),
        "Travel Misc": labels("Outras viagens", "Otros viajes", "Other travel"),
        "Uncategorized": labels("Sem categoria", "Sin categoría", "Uncategorized"),
        "Utilities": labels("Utilidades", "Servicios básicos", "Utilities"),
    }
    groups = defaultdict(lambda: {"amount": 0.0, "transactions": 0})
    raw_by_country_month = defaultdict(float)
    category_originals = defaultdict(set)
    skipped_text = []
    negative_transactions = 0
    negative_amount = 0
    country_lookup = {"SV_"+name: code for code, name, _ in COUNTRIES}
    for row_number, row in enumerate(raw[1:], start=2):
        department, amount, month_name = row[4], row[5], row[6]
        code = country_lookup[department]
        month = MONTHS.index(month_name) + 1
        original_category = str(row[3]).strip() if row[3] is not None else "Uncategorized"
        category = aliases.get(original_category, original_category)
        category_originals[category].add(original_category)
        if not num(amount):
            skipped_text.append({"row": row_number, "countryCode": code, "month": month,
                                 "amount": float(amount), "sourceCell": f"F{row_number}"})
            continue
        groups[(code, month, category)]["amount"] += amount
        groups[(code, month, category)]["transactions"] += 1
        raw_by_country_month[(code, month)] += amount
        if amount < 0:
            negative_transactions += 1
            negative_amount += amount
    card_records = [{"countryCode": code, "month": month, "category": category, **value}
                    for (code, month, category), value in sorted(groups.items())]
    card_deltas = []
    for record in records:
        delta = record["costs"]["card"] - raw_by_country_month[(record["countryCode"], record["month"])]
        card_deltas.append(delta)
        assert abs(delta) < .00001, (record["countryCode"], record["month"], delta)
    income_deltas = []
    for month in range(1, 10):
        auxiliary = expenses.cell(42+month, 2).value
        total = monthly[month-1]["income"]
        income_deltas.append({"month": month, "sourceCell": f"Expenses!B{42+month}",
                              "auxiliaryIncome": auxiliary, "pnlIncome": total, "difference": auxiliary-total})
    formula_errors = [{"sheet": sheet.title, "cell": cell.coordinate, "error": cell.value}
                      for sheet in values for row in sheet.iter_rows() for cell in row if cell.data_type == "e"]
    assert not formula_errors
    assert abs(totals["expenses"] - expenses["W25"].value) < .00001
    assert abs(totals["costs"]["card"] - expenses["K11"].value) < .00001

    issues = [
        {"id": "income_summary_difference", "severity": "warning", "status": "open", "amount": -652.50,
         "title": labels("Resumo de receitas divergente", "Resumen de ingresos divergente", "Income summary mismatch"),
         "detail": labels("O resumo auxiliar tem $652,50 a menos que a P&L. Abril e setembro explicam $651,96. O relatório usa a P&L por país e mês.",
                          "El resumen auxiliar tiene $652,50 menos que P&L. Abril y septiembre explican $651,96. El informe utiliza P&L por país y mes.",
                          "The auxiliary income summary is $652.50 below P&L. April and September account for $651.96. This report uses country-month P&L."),
         "source": {"sheet": "Expenses", "range": "A42:B52"}, "byMonth": income_deltas},
        {"id": "text_card_amounts", "severity": "warning", "status": "open", "amount": sum(row["amount"] for row in skipped_text),
         "count": len(skipped_text),
         "title": labels("Sete despesas fora do pivô", "Siete gastos fuera de la tabla dinámica", "Seven expenses excluded from the pivot"),
         "detail": labels("Sete despesas de junho, no total de $192,62, estão como texto e ficam fora do pivô e da P&L. Pendentes de revisão contábil.",
                          "Siete gastos de junio, por un total de $192,62, están guardados como texto y quedan fuera de la tabla dinámica y del P&L. Pendientes de revisión contable.",
                          "Seven June expenses totaling $192.62 are stored as text and excluded from the pivot and P&L. Pending accounting review."),
         "source": {"sheet": "Card Raw", "range": "A2:G31808"}, "excludedAmounts": skipped_text},
        {"id": "blank_colombia_billables", "severity": "info", "status": "documented", "amount": 0,
         "title": labels("Reembolsáveis da Colômbia em branco", "Reembolsables de Colombia en blanco", "Colombia billables cell is blank"),
         "detail": labels("Setembro tem célula vazia em Billables. A fórmula da fonte a trata como zero, com receita total igual à receita base de $3,68.",
                          "Septiembre tiene una celda vacía en Billables. La fórmula de origen la trata como cero y los ingresos totales igualan los ingresos base de $3,68.",
                          "September has a blank Billables cell. The source formula treats it as zero, with total revenue equal to base income of $3.68."),
         "source": {"sheet": "P&L", "range": "C52:F52"}},
        {"id": "presentation_allocation_scope", "severity": "info", "status": "explained", "amount": 17737.571,
         "title": labels("PPT: países antes do rateio", "PPT: países antes del reparto", "PPT: countries before cost allocation"),
         "detail": labels("Em setembro, os cinco cards de países do PPT excluem $2.351,28 de rateio por país e não mostram a Colômbia. O total LATAM inclui ambos. A ponte de despesas é $11.756,38 em rateios + $5.981,19 da Colômbia = $17.737,57.",
                          "En septiembre, las cinco tarjetas de países del PPT excluyen $2.351,28 de reparto por país y no muestran Colombia. LATAM incluye ambos. El puente de gastos es $11.756,38 en repartos + $5.981,19 de Colombia = $17.737,57.",
                          "In September, the five country cards in the PPT exclude $2,351.28 of allocated costs per country and omit Colombia. LATAM includes both. The expense bridge is $11,756.38 of allocations + $5,981.19 for Colombia = $17,737.57."),
         "source": {"sheet": "Expenses", "range": "J19:J30,AF11:AG11", "presentation": "September MBR - Oct 15.pptx", "slides": [15,16,17,18,19,20]}},
    ]
    result = {
        "schemaVersion": 1,
        "meta": {"title": "LATAM Financial Report", "year": year, "firstMonth": 1, "lastMonth": 9,
                 "currency": "USD", "currencySymbol": "$", "currencyStatus": "confirmed_by_user", "unit": "currency_units",
                 "currencyConfirmation": labels("Valores em USD, sem conversão cambial.",
                                                "Importes en USD, sin conversión de divisas.",
                                                "Values in USD, with no currency conversion."),
                 "sourceFile": SOURCE.name, "sourceType": "workbook_snapshot",
                 "sourceModifiedAt": values.properties.modified.isoformat(),
                 "extractedAt": datetime.now(timezone.utc).isoformat(), "sourceSha256": sha256(SOURCE.read_bytes()).hexdigest(),
                 "dateRange": {"from": min(dated).date().isoformat(), "to": max(dated).date().isoformat()},
                 "recordCount": len(records), "countryCount": len(COUNTRIES), "transactionCount": len(raw)-1,
                 "includedTransactionCount": len(raw)-1-len(skipped_text), "managementFeeRate": .078,
                 "financialsLive": False,
                 "basis": labels("Fechamento de janeiro a setembro de 2026, conforme os valores salvos na planilha.",
                                  "Cierre de enero a septiembre de 2026, según los valores guardados en la hoja.",
                                  "January–September 2026 close, based on saved workbook values.")},
        "countries": [{"code": code, "name": country} for code, country, _ in COUNTRIES],
        "months": [{"month": index+1, "period": f"{year}-{index+1:02d}", "label": label} for index, label in enumerate(MONTHS)],
        "records": records, "totals": totals, "monthly": monthly, "countryTotals": country_totals,
        "costCategories": [
            {"id": "card", "label": labels("Cartão corporativo", "Tarjeta corporativa", "Corporate card")},
            {"id": "payroll", "label": labels("Folha operacional", "Nómina operativa", "Operating payroll")},
            {"id": "pmSalary", "label": labels("Gestão de projeto", "Gestión del proyecto", "Project management")},
            {"id": "seatCharges", "label": labels("Postos de trabalho", "Puestos de trabajo", "Seat charges")},
            {"id": "procurement", "label": labels("Compras", "Compras", "Procurement")},
            {"id": "cashTransfers", "label": labels("Transferências", "Transferencias", "Cash transfers")},
            {"id": "otherIndirect", "label": labels("Outros indiretos", "Otros indirectos", "Other indirect costs")},
        ],
        "cardCategoryRecords": card_records,
        "cardCategoryLabels": category_labels,
        "cardCategoryAliases": {key: sorted(items) for key, items in sorted(category_originals.items())},
        "methodology": {
            "revenueFormula": "income + billables + managementFees", "profitFormula": "revenue - expenses",
            "marginFormula": "profit / revenue", "returnOnCostFormula": "profit / expenses",
            "allocatedOverheadFormula": "pmSalary + seatCharges + procurement + otherIndirect",
            "localExpensesFormula": "expenses - allocatedOverhead",
            "contributionProfitFormula": "revenue - localExpenses", "contributionMarginFormula": "contributionProfit / revenue",
            "managementFeesFormula": "billables * 7.8%", "costAllocation": "PM salary, seat charges, procurement and other indirect costs are allocated equally across six countries in the source workbook.",
            "costAllocationLabels": labels("Gestão, postos de trabalho, compras e outros indiretos são divididos igualmente entre os seis países na planilha.",
                                           "Gestión, puestos de trabajo, compras y otros indirectos se reparten igualmente entre los seis países en la hoja.",
                                           "Management, seat charges, procurement and other indirect costs are divided equally across six countries in the workbook."),
            "cardBasis": "Numeric amounts only, matching the supplied pivot and P&L. Credits/refunds remain negative. Category variants are normalized by the explicit alias map.",
            "blankComponents": missing_components,
            "sourceRanges": {"pnl": [f"P&L!B{start}:H{start+9}" for _, _, start in COUNTRIES],
                             "card": "Expenses!A4:K11", "payroll": "Expenses!A18:K30", "cashTransfers": "Expenses!R3:AA10",
                             "procurement": "Expenses!M3:O11", "otherIndirect": "Expenses!AD3:AH12",
                             "expenseControl": "Expenses!M18:W25", "cardTransactions": "Card Raw!A1:G31808"}},
        "reconciliation": {
            "checks": {"revenueAndProfitFormulasMatch": True, "costComponentsMatchPnl": True,
                       "countryTotalsMatchSource": True, "cardCategoriesMatchPivot": True, "transactionDatesMatchMonths": True,
                       "maxCostComponentDifference": max(abs(n) for n in component_checks),
                       "maxCardPivotDifference": max(abs(n) for n in card_deltas), "cachedFormulaErrorCount": len(formula_errors)},
            "numericCardTotal": totals["costs"]["card"], "excludedTextCardAmount": sum(row["amount"] for row in skipped_text),
            "allParsedCardAmounts": totals["costs"]["card"]+sum(row["amount"] for row in skipped_text),
            "negativeCardTransactionCount": negative_transactions, "negativeCardAmount": negative_amount,
            "auxiliaryIncomeTotal": expenses["B52"].value, "pnlIncomeTotal": totals["income"],
            "auxiliaryIncomeDifference": expenses["B52"].value-totals["income"], "issues": issues,
        },
    }
    OUTPUT.write_text(json.dumps(tidy(result), ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "records": len(records), "cardCategoryRecords": len(card_records),
                      "totals": tidy(totals), "checks": tidy(result["reconciliation"]["checks"]),
                      "issues": [issue["id"] for issue in issues]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
