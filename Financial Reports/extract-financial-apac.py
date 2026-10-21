"""Extract APAC country/month financial aggregates without editing the workbook.

Only aggregate amounts, categories and source coordinates are exported. Employee
names, merchants and individual transactions are never included in the JSON.
Run with the bundled Python runtime; openpyxl is used for read-only extraction.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "APAC (Jan-Sep).xlsx"
OUTPUT = HERE / "financial-apac-data.json"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September"]
COUNTRIES = [
    ("TH", "Thailand", 5, 4, 16, 25, 9),
    ("AU", "Australia", 18, 5, 17, 26, 6),
    ("TW", "Taiwan", 31, 6, 18, 27, 8),
    ("NZ", "New Zealand", 44, 7, 19, 28, 7),
]
FIELDS = ["income", "billables", "managementFees", "revenue", "expenses", "profit"]
COSTS = ["card", "payroll", "pmSalary", "seatCharges", "procurement", "cashTransfers", "otherIndirect", "sourceAdjustment"]
EPSILON = 0.0000001


def labels(pt, es, en):
    return {"pt": pt, "es": es, "en": en}


COUNTRY_LABELS = {
    "TH": labels("Tailândia", "Tailandia", "Thailand"),
    "AU": labels("Austrália", "Australia", "Australia"),
    "TW": labels("Taiwan", "Taiwán", "Taiwan"),
    "NZ": labels("Nova Zelândia", "Nueva Zelanda", "New Zealand"),
}
CATEGORY_LABELS = {
    "Advertising": labels("Publicidade", "Publicidad", "Advertising"),
    "Airlines": labels("Passagens aéreas", "Pasajes aéreos", "Airlines"),
    "Alcohol and Bars": labels("Bares e bebidas", "Bares y bebidas", "Alcohol and bars"),
    "Car Rental": labels("Aluguel de veículos", "Alquiler de vehículos", "Car rental"),
    "Car Services": labels("Serviços de veículos", "Servicios de vehículos", "Vehicle services"),
    "Education": labels("Educação", "Educación", "Education"),
    "Entertainment": labels("Entretenimento", "Entretenimiento", "Entertainment"),
    "Fees and Financial institutions": labels("Tarifas financeiras", "Comisiones financieras", "Financial fees"),
    "Fuel and Gas": labels("Combustível", "Combustible", "Fuel and gas"),
    "General Merchandise": labels("Mercadorias em geral", "Mercancías generales", "General merchandise"),
    "Government Services": labels("Serviços públicos", "Servicios gubernamentales", "Government services"),
    "Internet and Phone": labels("Internet e telefone", "Internet y teléfono", "Internet and phone"),
    "Lodging": labels("Hospedagem", "Alojamiento", "Lodging"),
    "Office": labels("Escritório", "Oficina", "Office"),
    "Office Supplies": labels("Materiais de escritório", "Material de oficina", "Office supplies"),
    "Office supplies and cleaning": labels("Escritório e limpeza", "Oficina y limpieza", "Office supplies and cleaning"),
    "Other": labels("Outros", "Otros", "Other"),
    "Parking": labels("Estacionamento", "Estacionamiento", "Parking"),
    "Professional Services": labels("Serviços profissionais", "Servicios profesionales", "Professional services"),
    "Restaurants": labels("Restaurantes", "Restaurantes", "Restaurants"),
    "SaaS / Software": labels("Software e SaaS", "Software y SaaS", "Software and SaaS"),
    "Shipping": labels("Envios", "Envíos", "Shipping"),
    "Supermarkets and Grocery Stores": labels("Supermercados", "Supermercados", "Supermarkets and groceries"),
    "Taxes and tax preparation": labels("Impostos", "Impuestos", "Taxes and tax preparation"),
    "Taxi and Rideshare": labels("Táxi e aplicativos", "Taxi y aplicaciones", "Taxi and rideshare"),
    "Toll": labels("Pedágios", "Peajes", "Tolls"),
    "Travel Misc": labels("Outras viagens", "Otros viajes", "Other travel"),
    "Uncategorized": labels("Sem categoria", "Sin categoría", "Uncategorized"),
    "Utilities": labels("Utilidades", "Servicios básicos", "Utilities"),
}
CATEGORY_ALIASES = {
    "Car services": "Car Services", "Phone and Internet": "Internet and Phone",
    "Office supplies": "Office Supplies", "Office and Supplies": "Office Supplies",
    "Taxi & Rideshare": "Taxi and Rideshare",
}


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def tidy(value):
    if isinstance(value, float):
        rounded = round(value, 8)
        return 0.0 if rounded == 0 else rounded
    if isinstance(value, dict):
        return {key: tidy(item) for key, item in value.items()}
    if isinstance(value, list):
        return [tidy(item) for item in value]
    return value


def close(actual, expected, label):
    assert numeric(actual) and numeric(expected), (label, "non-numeric value")
    difference = actual - expected
    assert abs(difference) < EPSILON, (label, actual, expected, difference)
    return difference


def aggregate(records):
    result = {field: sum(row[field] for row in records) for field in FIELDS}
    result["costs"] = {key: sum(row["costs"][key] for row in records) for key in COSTS}
    result["margin"] = result["profit"] / result["revenue"] if result["revenue"] else None
    result["returnOnCost"] = result["profit"] / result["expenses"] if result["expenses"] else None
    for field in ["allocatedOverhead", "localExpenses", "contributionProfit"]:
        result[field] = sum(row[field] for row in records)
    result["contributionMargin"] = result["contributionProfit"] / result["revenue"] if result["revenue"] else None
    kms = [row["billedKms"] for row in records if row["billedKms"] is not None]
    result["billedKms"] = sum(kms) if kms else None
    result["billedKmsMissingRecords"] = len(records) - len(kms)
    result["billedKmsCoverage"] = "missing" if not kms else "partial" if len(kms) < len(records) else "recorded"
    return result


def main():
    source_hash = sha256(SOURCE.read_bytes()).hexdigest()
    values = load_workbook(SOURCE, read_only=True, data_only=True)
    formulas = load_workbook(SOURCE, read_only=True, data_only=False)
    assert values.sheetnames == ["Expenses & Revenue", "P&L", "Card Raw"]
    pnl, expense = values["P&L"], values["Expenses & Revenue"]
    raw = list(values["Card Raw"].values)
    assert raw[0] == ("Transaction Date", "User", "Merchant Name", "Ramp Category", "Ramp Department", "Statement Amount", "Month")
    country_lookup = {"SV_" + name: code for code, name, *_ in COUNTRIES}
    dated, category_originals = [], defaultdict(set)
    cards = defaultdict(float)
    card_groups = defaultdict(lambda: {"amount": 0.0, "transactions": 0})
    date_mismatches = defaultdict(lambda: {"count": 0, "amount": 0.0, "storedDateMonths": Counter()})
    negative_count, negative_amount = 0, 0.0
    for row in raw[1:]:
        # Identity/merchant columns are never copied into output structures.
        date = row[0] if isinstance(row[0], datetime) else datetime.strptime(row[0], "%m/%d/%Y")
        dated.append(date)
        code = country_lookup[row[4]]
        month = MONTHS.index(row[6]) + 1
        amount = row[5]
        assert numeric(amount), "Unexpected non-numeric card amount; review source before publishing"
        original = str(row[3]) if row[3] is not None else "Uncategorized"
        category = CATEGORY_ALIASES.get(original.strip(), original.strip())
        assert category in CATEGORY_LABELS, ("Untranslated category", category)
        category_originals[category].add(original)
        cards[(code, month)] += amount
        group = card_groups[(code, month, category)]
        group["amount"] += amount
        group["transactions"] += 1
        if amount < 0:
            negative_count += 1
            negative_amount += amount
        if date.month != month:
            mismatch = date_mismatches[(code, month)]
            mismatch["count"] += 1
            mismatch["amount"] += amount
            mismatch["storedDateMonths"][date.month] += 1
    years = {date.year for date in dated}
    assert len(years) == 1, "Mixed years require an explicit reporting-period review"
    year = years.pop()
    assert expense["P3"].value == "For January only"
    assert expense["Q3"].value == "Cash Transfers (IDC)"
    assert expense["E22"].value == "Billed KMS"
    assert expense["E10"].value == "India Seat Charges"

    records, blank_components, adjustments, rates = [], [], [], set()
    country_differences, income_differences, kms_differences, payroll_differences = [], [], [], []
    component_differences, pivot_differences, unadjusted_differences = [], [], []
    for code, country, start, payroll_row, income_row, kms_row, pivot_row in COUNTRIES:
        assert pnl.cell(start - 1, 1).value == country
        assert [pnl.cell(start - 1, col).value for col in range(2, 10)] == [
            "Month", "KMS", "Income", "Billables (Repair & tech)", "Mgmt Fees", "Total Revenue", "Total Expenses", "P&L"]
        assert expense.cell(payroll_row, 5).value == "SV_" + country
        assert expense.cell(income_row, 2).value == "SV_" + country
        assert expense.cell(kms_row, 2).value == country
        assert expense.cell(pivot_row, 2).value == "SV_" + country
        for month, month_name in enumerate(MONTHS, 1):
            row = start + month - 1
            assert pnl.cell(row, 2).value == month_name
            assert expense.cell(3, month + 5).value == month_name
            assert expense.cell(15, month + 2).value == month_name
            assert expense.cell(24, month + 2).value == month_name
            record = {"country": country, "countryCode": code, "month": month, "period": f"{year}-{month:02d}"}
            for col, field in enumerate(FIELDS, 4):
                value = pnl.cell(row, col).value
                if value is None and field == "billables":
                    close(pnl.cell(row, 7).value, pnl.cell(row, 4).value, "blank billables revenue")
                    close(pnl.cell(row, 6).value, 0, "blank billables fee")
                    value = 0
                    blank_components.append({"countryCode": code, "month": month, "field": field,
                        "cell": f"P&L!E{row}", "treatment": "zero_as_used_by_source_formula"})
                assert numeric(value), (code, month, field)
                record[field] = value
            record["billedKms"] = pnl.cell(row, 3).value
            if record["billedKms"] is None:
                blank_components.append({"countryCode": code, "month": month, "field": "billedKms",
                    "cell": f"P&L!C{row}", "treatment": "preserved_as_missing"})
                assert expense.cell(kms_row, month + 2).value is None
            else:
                kms_differences.append(close(record["billedKms"], expense.cell(kms_row, month + 2).value, "billed KMS"))
            income_differences.append(close(record["income"], expense.cell(income_row, month + 2).value, "income"))
            fee_formula = formulas["P&L"].cell(row, 6).value
            match = re.fullmatch(rf"=E{row}\*([0-9.]+)%", fee_formula or "")
            assert match, ("Unexpected management fee formula", row)
            rate = float(match[1]) / 100
            rates.add(rate)
            close(record["managementFees"], record["billables"] * rate, "management fee")
            assert formulas["P&L"].cell(row, 7).value == f"=D{row}+E{row}+F{row}"
            assert formulas["P&L"].cell(row, 9).value == f"=G{row}-H{row}"
            close(record["revenue"], record["income"] + record["billables"] + record["managementFees"], "revenue")
            close(record["profit"], record["revenue"] - record["expenses"], "profit")
            payroll_ref = f"{get_column_letter(month + 5)}{payroll_row}"
            cash_ref = f"Q{payroll_row}" if month == 1 else None
            costs = dict.fromkeys(COSTS, 0)
            costs["card"] = cards[(code, month)]
            costs["payroll"] = expense[payroll_ref].value
            costs["cashTransfers"] = expense[cash_ref].value if cash_ref else 0
            assert all(numeric(value) for value in costs.values())
            delta = record["expenses"] - sum(costs.values())
            unadjusted_differences.append(delta)
            if abs(delta) >= EPSILON:
                # Explicit source difference, never mislabeled as an operating cost.
                # Larger/new differences require review instead of an automatic plug.
                assert code == "TW" and month == 7 and abs(delta + 0.05) < EPSILON, ("New unreconciled expense", code, month, delta)
                costs["sourceAdjustment"] = delta
                adjustments.append({"countryCode": code, "month": month, "amount": delta,
                    "pnlExpenses": record["expenses"], "card": costs["card"], "payroll": costs["payroll"],
                    "cashTransfers": costs["cashTransfers"], "pnlCell": f"P&L!H{row}",
                    "payrollCell": f"Expenses & Revenue!{payroll_ref}"})
            record["costs"] = costs
            component_differences.append(close(record["expenses"], sum(costs.values()), "costs including disclosed difference"))
            record["allocatedOverhead"] = 0
            record["localExpenses"] = record["expenses"]
            record["contributionProfit"] = record["profit"]
            record["margin"] = record["profit"] / record["revenue"] if record["revenue"] else None
            record["returnOnCost"] = record["profit"] / record["expenses"] if record["expenses"] else None
            record["contributionMargin"] = record["margin"]
            record["source"] = {"sheet": "P&L", "range": f"B{row}:I{row}", "billedKmsCell": f"C{row}",
                "costSheet": "Expenses & Revenue", "costCells": {"payroll": payroll_ref, "cashTransfers": cash_ref,
                    "card": f"C{pivot_row}" if month == 1 else None},
                "cardAggregation": {"sheet": "Card Raw", "range": f"D2:G{len(raw)}", "department": "SV_" + country, "month": month_name}}
            records.append(record)
        for col, field in enumerate(FIELDS, 4):
            country_differences.append(close(sum(r[field] for r in records if r["countryCode"] == code), pnl.cell(start + 9, col).value, "country P&L control"))
        close(sum(r["billedKms"] or 0 for r in records if r["countryCode"] == code), pnl.cell(start + 9, 3).value, "country KMS control")
        pivot_differences.append(close(cards[(code, 1)], expense.cell(pivot_row, 3).value, "January card pivot"))

    assert len(rates) == 1
    totals = aggregate(records)
    monthly = [{"month": month, "period": f"{year}-{month:02d}", **aggregate([r for r in records if r["month"] == month])} for month in range(1, 10)]
    country_totals = [{"countryCode": code, "country": name, **aggregate([r for r in records if r["countryCode"] == code])} for code, name, *_ in COUNTRIES]
    for row in monthly:
        month = row["month"]
        close(row["income"], expense.cell(20, month + 2).value, "monthly income control")
        close(row["billedKms"], expense.cell(29, month + 2).value, "monthly billed KMS control")
        payroll_differences.append(close(row["costs"]["payroll"], expense.cell(8, month + 5).value, "monthly payroll control"))
    close(totals["costs"]["cashTransfers"], expense["Q8"].value, "cash transfers control")
    close(monthly[0]["costs"]["card"], expense["C10"].value, "January total card pivot")
    errors = [{"sheet": sheet.title, "cell": cell.coordinate, "error": cell.value}
              for sheet in values for cells in sheet for cell in cells if cell.data_type == "e"]
    assert not errors, "Cached spreadsheet errors require review"
    seat_charge = expense["F10"].value
    assert formulas["Expenses & Revenue"]["F10"].value == "=1349.79*9"
    close(seat_charge, 1349.79 * 9, "standalone India seat charge")
    mismatches = [{"countryCode": code, "month": month, **item} for (code, month), item in sorted(date_mismatches.items())]
    mismatch_count = sum(item["count"] for item in mismatches)
    mismatch_amount = sum(item["amount"] for item in mismatches)
    allocation_labels = labels(
        "A P&L inclui cartão, folha e transferências de janeiro por país. Não há rateio de custos indiretos. Os postos da Índia aparecem separados e não integram a P&L. A diferença de conciliação de Taiwan está identificada à parte.",
        "P&L incluye tarjeta, nómina y transferencias de enero por país. No hay reparto de costes indirectos. Los puestos de India figuran por separado y no integran P&L. La diferencia de conciliación de Taiwán se identifica aparte.",
        "P&L includes country card costs, payroll and January transfers. No indirect-cost allocation is applied. India seat charges appear separately and are not included in P&L. Taiwan's reconciliation difference is identified separately.")
    issues = [
        {"id": "apac_india_seat_charges_outside_pnl", "severity": "warning", "status": "open", "amount": seat_charge,
         "title": labels("Postos da Índia fora da P&L", "Puestos de India fuera de P&L", "India seat charges outside P&L"),
         "detail": labels("A fonte registra US$ 12.148,11 em postos da Índia (US$ 1.349,79 × 9), sem distribuição por país. Esse valor não entra nas despesas da P&L. O relatório preserva a P&L e não inventa um rateio.",
                          "La fuente registra US$ 12.148,11 en puestos de India (US$ 1.349,79 × 9), sin distribución por país. Este importe no se incluye en los gastos de P&L. El informe conserva P&L y no inventa un reparto.",
                          "The source lists US$12,148.11 of India seat charges (US$1,349.79 × 9), without a country allocation. This amount is excluded from P&L expenses. The report preserves P&L and applies no invented allocation."),
         "source": {"sheet": "Expenses & Revenue", "range": "E10:F10"}},
        {"id": "apac_july_transaction_dates", "severity": "warning", "status": "open", "amount": mismatch_amount, "count": mismatch_count,
         "title": labels("Datas divergentes nos lançamentos de julho", "Fechas divergentes en los movimientos de julio", "Date mismatches in July transactions"),
         "detail": labels(f"{mismatch_count} lançamentos classificados como julho têm datas armazenadas em outros meses. As categorias seguem a coluna Month, que reconcilia com a P&L, sem mudar as datas originais.",
                          f"{mismatch_count} movimientos clasificados como julio tienen fechas guardadas en otros meses. Las categorías siguen la columna Month, que concilia con P&L, sin cambiar las fechas originales.",
                          f"{mismatch_count} transactions classified as July have stored dates in other months. Categories follow the Month column, which reconciles with P&L, without changing source dates."),
         "source": {"sheet": "Card Raw", "range": f"A2:G{len(raw)}"}, "byCountryMonth": mismatches},
        {"id": "apac_taiwan_cost_difference", "severity": "info", "status": "open", "amount": sum(item["amount"] for item in adjustments),
         "title": labels("Diferença de US$ 0,05 em Taiwan", "Diferencia de US$ 0,05 en Taiwán", "US$0.05 Taiwan difference"),
         "detail": labels("Em julho, a despesa da P&L de Taiwan é US$ 0,05 menor que cartão mais folha. A linha Diferença da fonte preserva o total publicado e identifica a diferença, sem atribuí-la a uma categoria operacional.",
                          "En julio, el gasto de P&L de Taiwán es US$ 0,05 menor que tarjeta más nómina. La línea Diferencia de origen conserva el total publicado e identifica la diferencia sin atribuirla a una categoría operativa.",
                          "Taiwan's July P&L expense is US$0.05 below card costs plus payroll. The Source difference line preserves the reported total and identifies the difference without assigning it to an operating category."),
         "source": {"sheet": "P&L", "range": "H37"}, "differences": adjustments},
        {"id": "apac_new_zealand_blank_components", "severity": "info", "status": "documented", "amount": 0,
         "title": labels("Nova Zelândia: campos vazios de julho a setembro", "Nueva Zelanda: campos vacíos de julio a septiembre", "New Zealand: blank July–September fields"),
         "detail": labels("A P&L registra receita zero de julho a setembro e despesa de US$ 24,57 em julho. Billables vazios entram como zero nas fórmulas da fonte. KMS vazios permanecem indisponíveis; não são convertidos em zero.",
                          "P&L registra ingresos cero de julio a septiembre y un gasto de US$ 24,57 en julio. Los Billables vacíos cuentan como cero en las fórmulas de origen. Los KMS vacíos siguen sin estar disponibles; no se convierten en cero.",
                          "P&L records zero revenue in July–September and US$24.57 of July expense. Blank Billables contribute zero in source formulas. Blank KMS remain unavailable and are not converted to zero."),
         "source": {"sheet": "P&L", "range": "B50:I52"}},
    ]
    result = {
        "schemaVersion": 1,
        "meta": {"title": "APAC Financial Report", "region": "apac", "reportRegionLabel": labels("Ásia-Pacífico", "Asia-Pacífico", "Asia-Pacific"),
            "year": year, "firstMonth": 1, "lastMonth": 9, "currency": "USD", "currencySymbol": "$", "currencyStatus": "confirmed_by_user", "unit": "currency_units",
            "currencyConfirmation": labels("Os valores financeiros da coleção são em USD, conforme confirmação do usuário. Nenhuma conversão cambial foi aplicada.", "Los valores financieros de la colección están en USD, según la confirmación del usuario. No se aplicó conversión de divisas.", "The collection's financial amounts are in USD, as confirmed by the user. No currency conversion was applied."),
            "sourceFile": SOURCE.name, "sourceType": "workbook_snapshot", "sourceModifiedAt": values.properties.modified.isoformat(),
            "extractedAt": datetime.now(timezone.utc).isoformat(), "sourceSha256": source_hash,
            "dateRange": {"from": f"{year}-01-01", "to": f"{year}-09-30"}, "dateRangeBasis": "P&L month labels and Card Raw Month column",
            "storedTransactionDateRange": {"from": min(dated).date().isoformat(), "to": max(dated).date().isoformat()},
            "recordCount": len(records), "countryCount": len(COUNTRIES), "transactionCount": len(raw) - 1,
            "includedTransactionCount": len(raw) - 1, "managementFeeRate": rates.pop(), "financialsLive": False,
            "operationsLive": False, "idleAvailable": False,
            "basis": labels("Valores salvos na planilha APAC. Períodos seguem a P&L e a coluna Month; todos os anos dos lançamentos são 2026. Divergências de datas estão documentadas.", "Valores guardados en la hoja APAC. Los períodos siguen P&L y la columna Month; todos los años de los movimientos son 2026. Las divergencias de fechas están documentadas.", "Cached APAC workbook values. Periods follow P&L and the Month column; all transaction years are 2026. Date mismatches are documented.")},
        "countries": [{"code": code, "name": name, "labels": COUNTRY_LABELS[code]} for code, name, *_ in COUNTRIES],
        "months": [{"month": month, "period": f"{year}-{month:02d}", "label": name} for month, name in enumerate(MONTHS, 1)],
        "records": records, "totals": totals, "monthly": monthly, "countryTotals": country_totals,
        "costCategories": [
            {"id": "card", "label": labels("Cartão corporativo", "Tarjeta corporativa", "Corporate card")},
            {"id": "payroll", "label": labels("Folha operacional", "Nómina operativa", "Operating payroll")},
            {"id": "pmSalary", "label": labels("Gestão de projeto", "Gestión del proyecto", "Project management")},
            {"id": "seatCharges", "label": labels("Postos de trabalho", "Puestos de trabajo", "Seat charges")},
            {"id": "procurement", "label": labels("Compras", "Compras", "Procurement")},
            {"id": "cashTransfers", "label": labels("Transferências", "Transferencias", "Cash transfers")},
            {"id": "otherIndirect", "label": labels("Outros indiretos", "Otros indirectos", "Other indirect costs")},
            {"id": "sourceAdjustment", "label": labels("Diferença da fonte", "Diferencia de origen", "Source difference")},
        ],
        "cardCategoryRecords": [{"countryCode": code, "month": month, "category": category, **group} for (code, month, category), group in sorted(card_groups.items())],
        "cardCategoryLabels": {category: CATEGORY_LABELS[category] for category in sorted(category_originals)},
        "cardCategoryAliases": {category: sorted(originals) for category, originals in sorted(category_originals.items())},
        "methodology": {
            "revenueFormula": "income + billables + managementFees", "profitFormula": "revenue - expenses", "marginFormula": "profit / revenue", "returnOnCostFormula": "profit / expenses",
            "allocatedOverheadFormula": "0 (no allocated overhead in source P&L)", "overheadCostKeys": [],
            "localExpensesFormula": "expenses - allocatedOverhead", "contributionProfitFormula": "revenue - localExpenses", "contributionMarginFormula": "contributionProfit / revenue",
            "managementFeesFormula": "billables * 7.8% (verified in each source formula)",
            "costAllocation": allocation_labels["en"], "costAllocationLabels": allocation_labels,
            "zeroCostKeys": ["pmSalary", "seatCharges", "procurement", "otherIndirect"],
            "zeroCostMeaning": labels("Zero significa não incluído na P&L fornecida; não comprova ausência do custo na empresa.", "Cero significa no incluido en P&L proporcionada; no demuestra ausencia del coste en la empresa.", "Zero means excluded from the supplied P&L; it does not establish that the company incurred no such cost."),
            "cardBasis": "All numeric Statement Amount rows, grouped by Ramp Department and the explicit Month column. Credits/refunds remain negative. No employee or merchant detail is exported.",
            "billedKmsBasis": labels("KMS da P&L reconciliados com a seção Billed KMS. É uma medida histórica da planilha financeira, sem sincronização operacional online.", "KMS de P&L conciliados con la sección Billed KMS. Es una medida histórica de la hoja financiera, sin sincronización operativa en línea.", "P&L KMS reconciled with the Billed KMS section. This is a historical financial-workbook measure, without live operational synchronization."),
            "blankComponents": blank_components,
            "sourceRanges": {"pnl": [f"P&L!B{start}:I{start + 9}" for _, _, start, *_ in COUNTRIES], "card": "Expenses & Revenue!B3:C10 (January only)",
                "payroll": "Expenses & Revenue!E3:N8", "cashTransfers": "Expenses & Revenue!P3:Q8 (January only)",
                "unallocatedSeatCharges": "Expenses & Revenue!E10:F10", "incomeControl": "Expenses & Revenue!B15:K20",
                "billedKmsControl": "Expenses & Revenue!B24:K29", "cardTransactions": f"Card Raw!A1:G{len(raw)}"}},
        "reconciliation": {
            "checks": {"revenueAndProfitFormulasMatch": True, "managementFeeFormulasVerified": True, "costComponentsMatchPnl": True,
                "unadjustedCostComponentsMatchPnl": not adjustments, "countryTotalsMatchSource": True, "cardCategoriesMatchPivot": True,
                "cardPivotCheckMonths": [1], "incomeMatchesAuxiliary": True, "payrollMatchesSource": True, "billedKmsMatchSource": True,
                "transactionDatesMatchMonths": not date_mismatches, "maxCostComponentDifference": max(map(abs, component_differences)),
                "maxUnadjustedCostDifference": max(map(abs, unadjusted_differences)), "maxCardPivotDifference": max(map(abs, pivot_differences)),
                "cachedFormulaErrorCount": len(errors)},
            "numericCardTotal": totals["costs"]["card"], "excludedTextCardAmount": 0, "allParsedCardAmounts": totals["costs"]["card"],
            "negativeCardTransactionCount": negative_count, "negativeCardAmount": negative_amount,
            "auxiliaryIncomeTotal": sum(expense.cell(20, month + 2).value for month in range(1, 10)), "pnlIncomeTotal": totals["income"],
            "auxiliaryIncomeDifference": 0, "sourceCostDifference": sum(item["amount"] for item in adjustments),
            "unallocatedSourceCosts": [{"id": "indiaSeatCharges", "amount": seat_charge, "source": {"sheet": "Expenses & Revenue", "range": "E10:F10"}}],
            "transactionDateMismatchCount": mismatch_count, "transactionDateMismatchAmount": mismatch_amount, "issues": issues},
    }
    values.close()
    formulas.close()
    assert sha256(SOURCE.read_bytes()).hexdigest() == source_hash, "Source workbook changed during extraction"
    OUTPUT.write_text(json.dumps(tidy(result), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "records": len(records), "countryCodes": [row[0] for row in COUNTRIES],
        "cardCategoryRecords": len(card_groups), "totals": tidy(totals), "checks": tidy(result["reconciliation"]["checks"]),
        "issues": [issue["id"] for issue in issues]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
