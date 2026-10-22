"""Extract the supplied Street View North America close without changing it.

Revenue, payroll and procurement are regional totals. Card transactions identify
US and Canada, but do not support allocating the full P&L between those countries.
The primary records therefore use the consolidated NA scope. Only country/month
card aggregates are exported separately. No employee, merchant or vendor names,
individual transactions, or Waymo data are published.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "SV-NA (Jan-Sep).xlsx"
OUTPUT = HERE / "financial-na-data.json"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September"]
FIELDS = ["income", "billables", "managementFees", "revenue", "expenses", "profit"]
COSTS = ["card", "payroll", "procurement", "otherIndirect", "nonBillablePayroll"]
COST_COLUMNS = {"card": "N", "payroll": "P", "procurement": "M", "otherIndirect": "O", "nonBillablePayroll": "Q"}
CARD_COUNTRIES = {"SV USA": "US", "SV Canada": "CA"}
TOLERANCE = .00001


def labels(pt, es, en):
    return {"pt": pt, "es": es, "en": en}


def num(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def tidy(value):
    if isinstance(value, float):
        return round(value, 8)
    if isinstance(value, dict):
        return {key: tidy(item) for key, item in value.items()}
    if isinstance(value, list):
        return [tidy(item) for item in value]
    return value


def check(actual, expected, context):
    difference = actual - expected
    if abs(difference) >= TOLERANCE:
        raise ValueError(f"Reconciliation failed: {context}: {difference}")
    return difference


def parse_card_date(value, month_name):
    if isinstance(value, datetime):
        return value
    candidates = set()
    for pattern in ["%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y", "%d/%m/%y", "%Y-%m-%d"]:
        try:
            candidate = datetime.strptime(str(value), pattern)
        except ValueError:
            continue
        if candidate.strftime("%B") == month_name:
            candidates.add(candidate)
    if len(candidates) != 1:
        raise ValueError("Ambiguous or invalid text transaction date")
    return next(iter(candidates))


def aggregate(rows):
    result = {key: sum(row[key] for row in rows) for key in FIELDS}
    result["costs"] = {key: sum(row["costs"][key] for row in rows) for key in COSTS}
    for key in ["allocatedOverhead", "localExpenses", "contributionProfit", "kms"]:
        result[key] = sum(row[key] for row in rows)
    result["margin"] = result["profit"] / result["revenue"] if result["revenue"] else None
    result["returnOnCost"] = result["profit"] / result["expenses"] if result["expenses"] else None
    result["contributionMargin"] = result["contributionProfit"] / result["revenue"] if result["revenue"] else None
    return result


def main():
    values = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
    try:
        pl, costs, payroll = values["Revenue"], values["IDC Expenses"], values["DC Expenses"]
        assert payroll["F3"].value == "SV - NA - 2026 - MoM Payroll"
        year = 2026  # Explicitly identified by the source payroll heading.
        assert pl["A2"].value == "Month" and pl["H2"].value == "P&L"
        assert formulas["IDC Expenses"]["H12"].value == "=H11/9"
        errors = [{"sheet": sheet.title, "cell": cell.coordinate, "error": cell.value}
                  for sheet in values for row in sheet.iter_rows() for cell in row if cell.data_type == "e"]
        if errors:
            raise ValueError("Source workbook contains cached spreadsheet errors")

        records, component_differences, income_differences = [], [], []
        for month, month_name in enumerate(MONTHS, 1):
            r, cr = month + 2, month + 3
            assert pl.cell(r, 1).value == month_name
            record = {"country": "North America (consolidated)", "countryCode": "NA", "month": month,
                      "period": f"{year}-{month:02d}", "kms": pl.cell(r, 2).value}
            for col, key in enumerate(FIELDS, 3):
                value = pl.cell(r, col).value
                assert num(value), (key, month)
                record[key] = value
            # The fee is explicit in Revenue E3:E11; it is not inherited from LATAM.
            assert formulas["Revenue"].cell(r, 5).value == f"=D{r}*7.8%"
            check(record["managementFees"], record["billables"] * .078, f"fee month {month}")
            check(record["revenue"], record["income"] + record["billables"] + record["managementFees"], f"revenue month {month}")
            check(record["profit"], record["revenue"] - record["expenses"], f"profit month {month}")
            refs = {key: f"{column}{cr}" for key, column in COST_COLUMNS.items()}
            record["costs"] = {key: costs[ref].value for key, ref in refs.items()}
            assert all(num(value) for value in record["costs"].values())
            component_differences.append(check(sum(record["costs"].values()), record["expenses"], f"cost components month {month}"))
            check(record["expenses"], costs[f"R{cr}"].value, f"expense close month {month}")
            check(record["costs"]["payroll"], payroll.cell(month+4, 7).value, f"monthly payroll {month}")
            check(record["costs"]["otherIndirect"], costs["H12"].value, f"monthly services {month}")
            # Only services have an explicit Jan-Sep amount / 9 in the source.
            # Procurement and both payroll lines retain their actual monthly totals.
            record["allocatedOverhead"] = record["costs"]["otherIndirect"]
            record["localExpenses"] = record["expenses"] - record["allocatedOverhead"]
            record["contributionProfit"] = record["revenue"] - record["localExpenses"]
            record["margin"] = record["profit"] / record["revenue"] if record["revenue"] else None
            record["returnOnCost"] = record["profit"] / record["expenses"] if record["expenses"] else None
            record["contributionMargin"] = record["contributionProfit"] / record["revenue"] if record["revenue"] else None
            record["source"] = {"sheet": "Revenue", "range": f"A{r}:H{r}", "costSheet": "IDC Expenses",
                                "costCells": refs, "kmsCell": f"B{r}"}
            records.append(record)
            income_differences.append({"month": month, "sourceCell": f"DC Expenses!H{month+4}",
                                       "auxiliaryIncome": payroll.cell(month+4, 8).value,
                                       "pnlIncome": record["income"],
                                       "difference": payroll.cell(month+4, 8).value-record["income"]})
        totals = aggregate(records)
        source_total_differences = [check(totals[key], pl.cell(12, col).value, f"Revenue total {key}")
                                    for col, key in enumerate(FIELDS, 3)]
        check(totals["kms"], pl["B12"].value, "source KMS total")
        check(totals["expenses"], costs["R13"].value, "IDC expense control")

        aliases = {"Fuel and gas": "Fuel and Gas", "Government services": "Government Services",
                   "Car services": "Car Services", "Phone and Internet": "Internet and Phone",
                   "Vehicle Maintanence": "Vehicle Maintenance"}
        category_labels = {
            "Airlines": labels("Passagens aéreas", "Pasajes aéreos", "Airlines"),
            "Alcohol and Bars": labels("Bares e bebidas", "Bares y bebidas", "Alcohol and bars"),
            "Books and Newspapers": labels("Livros e jornais", "Libros y periódicos", "Books and newspapers"),
            "Car Rental": labels("Aluguel de veículos", "Alquiler de vehículos", "Car rental"),
            "Car Services": labels("Serviços de veículos", "Servicios de vehículos", "Vehicle services"),
            "Education": labels("Educação", "Educación", "Education"),
            "Electronics": labels("Eletrônicos", "Electrónica", "Electronics"),
            "Entertainment": labels("Entretenimento", "Entretenimiento", "Entertainment"),
            "Freight, Moving and Delivery Services": labels("Fretes e entregas", "Fletes y entregas", "Freight and delivery"),
            "Fuel and Gas": labels("Combustível", "Combustible", "Fuel and gas"),
            "General Merchandise": labels("Mercadorias em geral", "Mercancías generales", "General merchandise"),
            "Government Services": labels("Serviços públicos", "Servicios gubernamentales", "Government services"),
            "Insurance": labels("Seguros", "Seguros", "Insurance"),
            "Internet and Phone": labels("Internet e telefone", "Internet y teléfono", "Internet and phone"),
            "Laundry": labels("Lavanderia", "Lavandería", "Laundry"),
            "Legal": labels("Jurídico", "Legal", "Legal"),
            "Lodging": labels("Hospedagem", "Alojamiento", "Lodging"),
            "Medical": labels("Saúde", "Salud", "Medical"),
            "Office supplies and cleaning": labels("Escritório e limpeza", "Oficina y limpieza", "Office supplies and cleaning"),
            "Other": labels("Outros", "Otros", "Other"),
            "Parking": labels("Estacionamento", "Estacionamiento", "Parking"),
            "Professional Services": labels("Serviços profissionais", "Servicios profesionales", "Professional services"),
            "Restaurants": labels("Restaurantes", "Restaurantes", "Restaurants"),
            "SaaS / Software": labels("Software e SaaS", "Software y SaaS", "Software and SaaS"),
            "Shipping": labels("Envios", "Envíos", "Shipping"),
            "Storage": labels("Armazenamento", "Almacenamiento", "Storage"),
            "Supermarkets and Grocery Stores": labels("Supermercados", "Supermercados", "Supermarkets and groceries"),
            "Taxi and Rideshare": labels("Táxi e aplicativos", "Taxi y aplicaciones", "Taxi and rideshare"),
            "Travel Misc": labels("Outras viagens", "Otros viajes", "Other travel"),
            "Vehicle Maintenance": labels("Manutenção de veículos", "Mantenimiento de vehículos", "Vehicle maintenance"),
        }
        card_rows = list(values["Card Raw"].iter_rows(values_only=True))
        assert card_rows[0] == ("Transaction Date", "User", "Merchant Name", "Ramp Category", "Ramp Department", "Statement Amount", "Month")
        category_groups = defaultdict(lambda: {"amount": 0.0, "transactions": 0})
        country_groups = defaultdict(lambda: {"amount": 0.0, "transactions": 0})
        country_category_groups = defaultdict(lambda: {"amount": 0.0, "transactions": 0})
        category_originals = defaultdict(set)
        date_month_mismatches = defaultdict(lambda: {"count": 0, "amount": 0.0})
        excluded_text = Counter()
        excluded_text_amount = 0.0
        dated, out_of_year, negative_count, negative_amount = [], [], 0, 0.0
        for rn, row in enumerate(card_rows[1:], 2):
            transaction_date, _, _, original_category, department, amount, month_name = row
            if department not in CARD_COUNTRIES or month_name not in MONTHS:
                raise ValueError("Unexpected card scope: review before publishing")
            country_code, month = CARD_COUNTRIES[department], MONTHS.index(month_name)+1
            dt = parse_card_date(transaction_date, month_name)
            dated.append(dt)
            if dt.year != year:
                out_of_year.append({"sourceCell": f"A{rn}", "recordedYear": dt.year, "sourceMonth": month})
            if dt.strftime("%B") != month_name:
                entry = date_month_mismatches[(dt.strftime("%Y-%m"), month)]
                entry["count"] += 1
                if num(amount):
                    entry["amount"] += amount
            if not num(amount):
                excluded_text[(country_code, month)] += 1
                excluded_text_amount += float(amount)
                continue  # The supplied numeric pivot excludes this text amount.
            category = aliases.get(original_category, original_category)
            assert category in category_labels, category
            category_originals[category].add(original_category)
            for group in [category_groups[(month, category)], country_groups[(country_code, month)],
                          country_category_groups[(country_code, month, category)]]:
                group["amount"] += amount
                group["transactions"] += 1
            if amount < 0:
                negative_count += 1
                negative_amount += amount
        card_records = [{"countryCode": "NA", "month": month, "category": category, **value}
                        for (month, category), value in sorted(category_groups.items())]
        card_country_records = [{"countryCode": code, "month": month, **country_groups[(code, month)]}
                                for code in ["US", "CA"] for month in range(1, 10)]
        card_country_category_records = [
            {"countryCode": code, "month": month, "category": category, **value}
            for (code, month, category), value in sorted(country_category_groups.items())]
        for row in card_country_records:
            detail = [c for c in card_country_category_records
                      if c["countryCode"] == row["countryCode"] and c["month"] == row["month"]]
            check(sum(c["amount"] for c in detail), row["amount"],
                  f"card country categories {row['countryCode']} month {row['month']}")
            assert sum(c["transactions"] for c in detail) == row["transactions"]
        for row in card_records:
            detail = [c for c in card_country_category_records
                      if c["month"] == row["month"] and c["category"] == row["category"]]
            check(sum(c["amount"] for c in detail), row["amount"],
                  f"card regional category {row['category']} month {row['month']}")
            assert sum(c["transactions"] for c in detail) == row["transactions"]
        card_differences = []
        for row in records:
            card_differences.append(check(sum(c["amount"] for c in card_records if c["month"] == row["month"]), row["costs"]["card"], f"card month {row['month']}"))
            check(sum(c["amount"] for c in card_country_records if c["month"] == row["month"]), row["costs"]["card"], f"card country month {row['month']}")
        check(totals["costs"]["card"], costs["E38"].value, "card pivot total")
        check(totals["costs"]["card"], values["Sheet6"]["B36"].value, "second card pivot total")
        check(totals["costs"]["card"], values["PPT Slides"]["H23"].value, "PPT card summary")

        procurement_rows = list(values["Proc Raw"].iter_rows(values_only=True))
        assert procurement_rows[0] == ("Month", "Category", "Project", "Vendor", "Actual Spent")
        procurement_groups = defaultdict(lambda: {"amount": 0.0, "transactions": 0})
        procurement_labels = {
            "Depot Consumables": labels("Consumíveis de depósito", "Consumibles de depósito", "Depot consumables"),
            "Flight": labels("Passagens aéreas", "Pasajes aéreos", "Flights"),
            "Hotel": labels("Hospedagem", "Alojamiento", "Hotels"),
            "Parking": labels("Estacionamento", "Estacionamiento", "Parking"),
            "Rental": labels("Locação", "Alquiler", "Rental"),
            "Services": labels("Serviços", "Servicios", "Services"),
            "Train": labels("Trem", "Tren", "Train"),
            "Travel Desk Serivce Fee": labels("Taxa de agência de viagens", "Tarifa de agencia de viajes", "Travel desk service fee"),
        }
        for month_name, category, project, _, amount in procurement_rows[1:]:
            if project != "SV-NA" or month_name not in MONTHS or not num(amount):
                raise ValueError("Unexpected procurement scope or amount")
            assert category in procurement_labels
            entry = procurement_groups[(MONTHS.index(month_name)+1, category)]
            entry["amount"] += amount
            entry["transactions"] += 1
        procurement_records = [{"countryCode": "NA", "month": month, "category": category, **value}
                               for (month, category), value in sorted(procurement_groups.items())]
        procurement_differences = [check(sum(p["amount"] for p in procurement_records if p["month"] == row["month"]), row["costs"]["procurement"], f"procurement month {row['month']}") for row in records]
        check(totals["costs"]["procurement"], values["PPT Slides"]["K13"].value, "PPT procurement summary")
        check(totals["costs"]["otherIndirect"], costs["H11"].value, "other services total")
        ppt_differences = []
        for row in records:
            rr = row["month"] + 3
            for column, key in [(18,"kms"),(19,"income"),(20,"billables"),(21,"revenue"),(22,"expenses"),(23,"profit")]:
                ppt_differences.append(check(row[key], values["PPT Slides"].cell(rr,column).value, f"PPT close month {row['month']} {key}"))

        auxiliary_income_delta = payroll["H17"].value - totals["income"]
        biweekly_payroll_delta = payroll["B33"].value - totals["costs"]["payroll"]
        payroll_summary_delta = payroll["G17"].value - totals["costs"]["payroll"]
        support_delta = payroll["L6"].value - totals["costs"]["nonBillablePayroll"]
        mismatch_count = sum(v["count"] for v in date_month_mismatches.values())
        issues = [
            {"id": "na_consolidated_scope", "severity": "info", "status": "documented",
             "title": labels("Resultado regional consolidado", "Resultado regional consolidado", "Consolidated regional result"),
             "detail": labels("A P&L reúne Estados Unidos e Canadá. A planilha separa apenas as despesas de cartão por país; receita, folha e demais custos são consolidados.", "P&L reúne Estados Unidos y Canadá. La hoja separa por país únicamente los gastos de tarjeta; los ingresos, la nómina y los demás gastos están consolidados.", "P&L combines the United States and Canada. Only card expenses are split by country; revenue, payroll and other costs are consolidated."),
             "source": {"sheet": "Revenue", "range": "A2:H12"}},
            {"id": "na_card_period_labels", "severity": "warning", "status": "open", "count": mismatch_count,
             "title": labels("Datas de cartão divergentes", "Fechas de tarjeta discrepantes", "Card date discrepancies"),
             "detail": labels(f"{mismatch_count} lançamentos têm datas fora do mês informado em Month; {len(out_of_year)} registra outro ano. O fechamento segue Month, conforme a tabela dinâmica e a P&L.", f"{mismatch_count} movimientos tienen fechas fuera del mes indicado en Month; {len(out_of_year)} corresponde a otro año. El cierre sigue Month, según la tabla dinámica y P&L.", f"{mismatch_count} card dates fall outside the month in the Month column; {len(out_of_year)} records another year. The close follows Month, consistent with the pivot table and P&L."),
             "source": {"sheet": "Card Raw", "range": "A2:G12605"},
             "bySourcePeriod": [{"transactionMonth": period, "sourceMonth": month, **entry} for (period,month),entry in sorted(date_month_mismatches.items())], "outOfYearCount": len(out_of_year), "outOfYearReferences": out_of_year},
            {"id": "na_text_card_amount", "severity": "warning", "status": "open", "count": sum(excluded_text.values()), "amount": excluded_text_amount,
             "title": labels("Valor de cartão armazenado como texto", "Importe de tarjeta guardado como texto", "Card amount stored as text"),
             "detail": labels("Um lançamento de USD 37,02 do Canadá (julho) está armazenado como texto e não entra na tabela dinâmica nem nas despesas do fechamento.", "Un movimiento de USD 37,02 de Canadá (julio) está guardado como texto y queda fuera de la tabla dinámica y de los gastos del cierre.", "A USD 37.02 Canada entry for July is stored as text and excluded from both the pivot table and closing expenses."),
             "source": {"sheet": "Card Raw", "range": "F9640"}},
            {"id": "na_auxiliary_income_difference", "severity": "warning", "status": "open", "amount": auxiliary_income_delta,
             "title": labels("Divergência de faturamento", "Diferencia de facturación", "Invoicing difference"),
             "detail": labels("Invoiced (DC Expenses) soma USD 954,19 a menos que Income (Revenue): USD 954,22 em abril, com diferenças menores em janeiro e agosto. Este relatório segue a aba Revenue, consistente com PPT Slides.", "Invoiced (DC Expenses) suma USD 954,19 menos que Income (Revenue): USD 954,22 en abril, con diferencias menores en enero y agosto. Este informe sigue la pestaña Revenue, consistente con PPT Slides.", "Invoiced (DC Expenses) totals USD 954.19 less than Income (Revenue): USD 954.22 in April, with smaller differences in January and August. This report follows the Revenue sheet, consistent with PPT Slides."),
             "source": {"sheet": "DC Expenses", "range": "F5:H17"}, "byMonth": income_differences},
            {"id": "na_payroll_auxiliary_difference", "severity": "warning", "status": "open", "amount": biweekly_payroll_delta,
             "title": labels("Divergência entre folhas quinzenal e mensal", "Diferencia entre nóminas quincenal y mensual", "Biweekly and monthly payroll differ"),
             "detail": labels("A P&L usa a folha mensal de USD 1.024.088,59. A listagem quinzenal soma USD 1.106.358,80, diferença de USD 82.270,21.", "P&L usa la nómina mensual de USD 1.024.088,59. La lista quincenal suma USD 1.106.358,80, una diferencia de USD 82.270,21.", "P&L uses monthly payroll of USD 1,024,088.59. The biweekly listing totals USD 1,106,358.80, a difference of USD 82,270.21."),
             "source": {"sheet": "DC Expenses", "range": "A3:I33"}},
            {"id": "na_payroll_rounding", "severity": "info", "status": "documented", "amount": support_delta,
             "title": labels("Diferenças de centavos na folha", "Diferencias de centavos en nómina", "Minor payroll total differences"),
             "detail": labels("Paid supera a soma mensal em USD 0,01. Em Other NB DC, o resumo soma USD 123.298,29 e as linhas mensais, USD 123.298,32. A P&L usa as somas mensais.", "Paid supera la suma mensual en USD 0,01. En Other NB DC, el resumen suma USD 123.298,29 y las líneas mensuales, USD 123.298,32. P&L usa las sumas mensuales.", "Paid exceeds the monthly sum by USD 0.01. Other NB DC totals USD 123,298.29 in the summary and USD 123,298.32 across monthly lines. P&L uses the monthly sums."),
             "source": {"sheet": "DC Expenses", "range": "G17,L4:L6"}, "payrollSummaryDifference": payroll_summary_delta},
        ]
        scope_countries = [
            {"code": "US", "name": "United States", "label": labels("Estados Unidos", "Estados Unidos", "United States")},
            {"code": "CA", "name": "Canada", "label": labels("Canadá", "Canadá", "Canada")},
        ]
        region_country = {"code": "NA", "name": "North America (consolidated)",
                          "label": labels("América do Norte (consolidado)", "Norteamérica (consolidado)", "North America (consolidated)")}
        result = {
            "schemaVersion": 1,
            "meta": {"title": "North America Financial Report", "region": "na", "year": year, "firstMonth": 1, "lastMonth": 9,
                     "currency": "USD", "currencySymbol": "$", "currencyStatus": "confirmed_by_user", "unit": "currency_units",
                     "currencyConfirmation": labels("Valores em USD, sem conversão cambial.", "Importes en USD, sin conversión de divisas.", "Amounts in USD; no currency conversion."),
                     "sourceFile": SOURCE.name, "sourceType": "workbook_snapshot", "pnlSheet": "Revenue",
                     "sourceModifiedAt": values.properties.modified.isoformat(), "extractedAt": datetime.now(timezone.utc).isoformat(),
                     "sourceSha256": sha256(SOURCE.read_bytes()).hexdigest(), "dateRange": {"from": "2026-01-01", "to": "2026-09-30"},
                     "dateRangeBasis": "Source Month column; transaction-date discrepancies are listed in reconciliation.",
                     "recordCount": len(records), "countryCount": 1, "scopeCountries": scope_countries, "countryBreakdown": "card-only",
                     "scopeNote": labels("Receitas e despesas consolidadas de Estados Unidos e Canadá. Apenas os gastos de cartão têm abertura por país.", "Ingresos y gastos consolidados de Estados Unidos y Canadá. Solo los gastos de tarjeta tienen desglose por país.", "Consolidated revenue and expenses for the United States and Canada. Only card spending is broken down by country."),
                     "transactionCount": len(card_rows)-1, "includedTransactionCount": sum(c["transactions"] for c in card_records),
                     "procurementTransactionCount": len(procurement_rows)-1, "managementFeeRate": .078, "financialsLive": False,
                     "basis": labels("Fechamento SV-NA de janeiro a setembro de 2026. Valores da aba Revenue; meses conforme a coluna Month.", "Cierre SV-NA de enero a septiembre de 2026. Importes de la pestaña Revenue; meses según la columna Month.", "SV-NA close for January–September 2026. Values from the Revenue sheet; periods follow the Month column.")},
            "countries": [region_country],
            "months": [{"month": m, "period": f"{year}-{m:02d}", "label": name} for m,name in enumerate(MONTHS,1)],
            "records": records, "totals": totals,
            "monthly": [{"month": row["month"], "period": row["period"], **aggregate([row])} for row in records],
            "countryTotals": [{"countryCode": "NA", "country": region_country["name"], **totals}],
            "costCategories": [
                {"id": "card", "label": labels("Cartão corporativo", "Tarjeta corporativa", "Corporate card")},
                {"id": "payroll", "label": labels("Folha operacional", "Nómina operativa", "Operating payroll")},
                {"id": "procurement", "label": labels("Compras e reservas", "Compras y reservas", "Procurement and bookings")},
                {"id": "otherIndirect", "label": labels("Outros serviços indiretos", "Otros servicios indirectos", "Other indirect services")},
                {"id": "nonBillablePayroll", "label": labels("Outros custos diretos não faturáveis", "Otros costos directos no facturables", "Other non-billable direct costs")},
            ],
            "cardCategoryRecords": card_records, "cardCategoryLabels": category_labels,
            "cardCategoryAliases": {key: sorted(items) for key,items in sorted(category_originals.items())},
            "cardCountryRecords": card_country_records,
            "cardCountryCategoryRecords": card_country_category_records,
            "procurementCategoryRecords": procurement_records, "procurementCategoryLabels": procurement_labels,
            "methodology": {
                "revenueFormula": "income + billables + managementFees", "profitFormula": "revenue - expenses",
                "marginFormula": "profit / revenue", "returnOnCostFormula": "profit / expenses", "managementFeesFormula": "billables * 7.8% (Revenue!E3:E11)",
                "overheadCostKeys": ["otherIndirect"], "allocatedOverheadFormula": "otherIndirect",
                "localExpensesFormula": "expenses - otherIndirect", "contributionProfitFormula": "revenue - localExpenses",
                "contributionMarginFormula": "contributionProfit / revenue",
                "costAllocation": "Other indirect services are allocated equally across January–September (IDC Expenses!H12 = H11/9). Procurement and both direct-cost lines use source monthly amounts, with no country allocation.",
                "costAllocationLabels": labels("Outros serviços indiretos rateados igualmente entre janeiro e setembro (IDC Expenses!H12 = H11/9). Compras e as duas linhas de custos diretos seguem os valores mensais, sem rateio entre países.", "Otros servicios indirectos repartidos por igual entre enero y septiembre (IDC Expenses!H12 = H11/9). Compras y ambas líneas de costos directos siguen los importes mensuales, sin reparto por país.", "Other indirect services are allocated equally across January–September (IDC Expenses!H12 = H11/9). Procurement and both direct-cost lines use monthly amounts, with no country allocation."),
                "cardBasis": "Numeric Statement Amount values by Month, including negative credits. Country totals use Ramp Department SV USA and SV Canada and cover card spending only.",
                "cardCountryLabels": labels("Cartão corporativo por país. Receita, folha e demais custos permanecem consolidados.", "Tarjetas corporativas por país. Los ingresos, la nómina y los demás gastos permanecen consolidados.", "Corporate card spending by country. Revenue, payroll and other costs remain consolidated."),
                "blankComponents": [],
                "sourceRanges": {"pnl": ["Revenue!A3:H12"], "expenseControl": "IDC Expenses!L3:R13", "cardTransactions": "Card Raw!A1:G12605", "procurementTransactions": "Proc Raw!A1:E936", "payroll": "DC Expenses!F4:I17", "otherIndirect": "IDC Expenses!G3:I12", "presentationControl": "PPT Slides!Q3:W13"}},
            "reconciliation": {
                "checks": {"revenueAndProfitFormulasMatch": True, "costComponentsMatchPnl": True, "countryTotalsMatchSource": True,
                           "cardCategoriesMatchPivot": True, "cardCountryTotalsMatchPivot": True, "procurementCategoriesMatchPnl": True,
                           "cardCountryCategoriesMatchCountryTotals": True,
                           "cardCountryCategoriesMatchRegionalCategories": True,
                           "presentationTotalsMatchSource": True, "transactionDatesMatchMonths": mismatch_count == 0,
                           "allTransactionYearsMatchReport": len(out_of_year) == 0, "cachedFormulaErrorCount": len(errors),
                           "maxCostComponentDifference": max(abs(x) for x in component_differences),
                           "maxCardPivotDifference": max(abs(x) for x in card_differences),
                           "maxProcurementDifference": max(abs(x) for x in procurement_differences),
                           "maxSourceTotalDifference": max(abs(x) for x in source_total_differences),
                           "maxPresentationDifference": max(abs(x) for x in ppt_differences)},
                "numericCardTotal": totals["costs"]["card"], "excludedTextCardAmount": excluded_text_amount,
                "allParsedCardAmounts": totals["costs"]["card"] + excluded_text_amount,
                "negativeCardTransactionCount": negative_count, "negativeCardAmount": negative_amount,
                "auxiliaryIncomeTotal": payroll["H17"].value, "pnlIncomeTotal": totals["income"], "auxiliaryIncomeDifference": auxiliary_income_delta,
                "biweeklyPayrollTotal": payroll["B33"].value, "monthlyPayrollTotal": totals["costs"]["payroll"], "biweeklyPayrollDifference": biweekly_payroll_delta,
                "nonBillableSummaryDifference": support_delta, "issues": issues},
        }
        OUTPUT.write_text(json.dumps(tidy(result), ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
        print(json.dumps({"output": OUTPUT.name, "records": len(records), "countryScopes": ["NA"],
                          "cardCountryRecords": len(card_country_records), "cardCategoryRecords": len(card_records),
                          "cardCountryCategoryRecords": len(card_country_category_records),
                          "procurementCategoryRecords": len(procurement_records), "totals": tidy(totals),
                          "checks": tidy(result["reconciliation"]["checks"]), "issues": [i["id"] for i in issues]}, ensure_ascii=True, indent=2))
    finally:
        values.close()
        formulas.close()


if __name__ == "__main__":
    main()
